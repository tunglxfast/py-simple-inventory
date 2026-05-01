"""Main business services for warehouse operations."""

from __future__ import annotations

import csv
import shutil
import sqlite3
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any
from uuid import uuid4

from database.connection import DatabaseManager
from services.exceptions import InsufficientStockError, NotFoundError, ValidationError


def utc_now_iso() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


@dataclass
class HoldItemInput:
    sku: str
    quantity: int


class WarehouseService:
    def __init__(self, db: DatabaseManager) -> None:
        self.db = db
        self.db.init_db()

    def create_product(self, sku: str, name: str, unit: str) -> None:
        sku = sku.strip()
        name = name.strip()
        unit = unit.strip()
        if not sku or not name or not unit:
            raise ValidationError("sku, name, unit are required")
        with self.db.transaction() as conn:
            conn.execute(
                "INSERT INTO products (sku, name, unit, quantity) VALUES (?, ?, ?, 0)",
                (sku, name, unit),
            )

    def update_product(self, sku: str, name: str, unit: str) -> None:
        with self.db.transaction() as conn:
            cur = conn.execute(
                "UPDATE products SET name = ?, unit = ? WHERE sku = ?",
                (name.strip(), unit.strip(), sku.strip()),
            )
            if cur.rowcount == 0:
                raise NotFoundError(f"Product not found: {sku}")

    def get_inventory(self) -> list[dict[str, Any]]:
        with self.db.connection() as conn:
            rows = conn.execute(
                "SELECT sku, name, unit, quantity FROM products ORDER BY sku"
            ).fetchall()
        return [dict(row) for row in rows]

    def import_stock(self, sku: str, quantity: int, note: str | None = None) -> None:
        self._validate_positive_qty(quantity)
        with self.db.transaction() as conn:
            self._apply_stock_change(
                conn=conn,
                sku=sku,
                change=quantity,
                movement_type="IMPORT",
                ref_type="IMPORT",
                ref_id=None,
            )

    def export_stock(self, sku: str, quantity: int, note: str | None = None) -> str:
        self._validate_positive_qty(quantity)
        with self.db.transaction() as conn:
            sale_code = self._new_code("SAL")
            now = utc_now_iso()
            sale_cur = conn.execute(
                "INSERT INTO sales (code, hold_session_id, created_at, note) VALUES (?, NULL, ?, ?)",
                (sale_code, now, note),
            )
            sale_id = int(sale_cur.lastrowid)
            conn.execute(
                "INSERT INTO sale_items (sale_id, sku, quantity) VALUES (?, ?, ?)",
                (sale_id, sku, quantity),
            )
            self._apply_stock_change(
                conn=conn,
                sku=sku,
                change=-quantity,
                movement_type="EXPORT",
                ref_type="EXPORT",
                ref_id=sale_id,
            )
        return sale_code

    def create_hold(self, items: list[dict[str, Any]], note: str | None = None) -> str:
        normalized_items = self._normalize_hold_items(items)
        with self.db.transaction() as conn:
            code = self._new_code("HLD")
            now = utc_now_iso()
            cur = conn.execute(
                "INSERT INTO hold_sessions (code, status, created_at, note) VALUES (?, 'OPEN', ?, ?)",
                (code, now, note),
            )
            session_id = int(cur.lastrowid)
            for item in normalized_items:
                conn.execute(
                    """
                    INSERT INTO hold_items (session_id, sku, quantity, returned_quantity, status)
                    VALUES (?, ?, ?, 0, 'HOLD')
                    """,
                    (session_id, item.sku, item.quantity),
                )
                self._apply_stock_change(
                    conn=conn,
                    sku=item.sku,
                    change=-item.quantity,
                    movement_type="EXPORT",
                    ref_type="HOLD",
                    ref_id=session_id,
                )
        return code

    def get_hold_sessions(self, query: str | None = None) -> list[dict[str, Any]]:
        sql = "SELECT id, code, status, created_at, note FROM hold_sessions"
        params: tuple[Any, ...] = ()
        if query:
            sql += " WHERE code LIKE ?"
            params = (f"%{query}%",)
        sql += " ORDER BY id DESC"
        with self.db.connection() as conn:
            rows = conn.execute(sql, params).fetchall()
        return [dict(row) for row in rows]

    def get_hold_details(self, code: str) -> dict[str, Any]:
        with self.db.connection() as conn:
            session = conn.execute(
                "SELECT id, code, status, created_at, note FROM hold_sessions WHERE code = ?",
                (code,),
            ).fetchone()
            if session is None:
                raise NotFoundError(f"Hold session not found: {code}")
            items = conn.execute(
                """
                SELECT hi.id, hi.sku, p.name, p.unit, hi.quantity, hi.returned_quantity, hi.status
                FROM hold_items hi
                JOIN products p ON p.sku = hi.sku
                WHERE hi.session_id = ?
                ORDER BY hi.id
                """,
                (session["id"],),
            ).fetchall()
        return {"session": dict(session), "items": [dict(row) for row in items]}

    def finalize_hold(self, session_code: str, returns: list[dict[str, Any]], note: str | None = None) -> str:
        returns_map = {entry["sku"]: int(entry["quantity"]) for entry in returns}
        with self.db.transaction() as conn:
            session = conn.execute(
                "SELECT id, status FROM hold_sessions WHERE code = ?", (session_code,)
            ).fetchone()
            if session is None:
                raise NotFoundError(f"Hold session not found: {session_code}")
            if session["status"] != "OPEN":
                raise ValidationError("Hold session is already finalized")

            items = conn.execute(
                "SELECT id, sku, quantity, returned_quantity FROM hold_items WHERE session_id = ?",
                (session["id"],),
            ).fetchall()
            sold_items: list[tuple[str, int]] = []
            for item in items:
                sku = item["sku"]
                held_qty = int(item["quantity"])
                return_qty = int(returns_map.get(sku, 0))
                if return_qty < 0 or return_qty > held_qty:
                    raise ValidationError(f"Invalid return qty for {sku}")
                if return_qty > 0:
                    self._apply_stock_change(
                        conn=conn,
                        sku=sku,
                        change=return_qty,
                        movement_type="IMPORT",
                        ref_type="RETURN",
                        ref_id=int(session["id"]),
                    )
                sold_qty = held_qty - return_qty
                status = "RETURNED" if return_qty == held_qty else "HOLD"
                conn.execute(
                    """
                    UPDATE hold_items
                    SET returned_quantity = ?, status = ?
                    WHERE id = ?
                    """,
                    (return_qty, status, item["id"]),
                )
                if sold_qty > 0:
                    sold_items.append((sku, sold_qty))

            sale_code = self._new_code("SAL")
            now = utc_now_iso()
            sale_cur = conn.execute(
                "INSERT INTO sales (code, hold_session_id, created_at, note) VALUES (?, ?, ?, ?)",
                (sale_code, session["id"], now, note),
            )
            sale_id = int(sale_cur.lastrowid)
            for sku, qty in sold_items:
                conn.execute(
                    "INSERT INTO sale_items (sale_id, sku, quantity) VALUES (?, ?, ?)",
                    (sale_id, sku, qty),
                )
            conn.execute(
                "UPDATE hold_sessions SET status = 'DONE' WHERE id = ?", (session["id"],)
            )
        return sale_code

    def get_sales(self) -> list[dict[str, Any]]:
        with self.db.connection() as conn:
            sales = conn.execute(
                "SELECT id, code, hold_session_id, created_at, note FROM sales ORDER BY id DESC"
            ).fetchall()
            output: list[dict[str, Any]] = []
            for sale in sales:
                items = conn.execute(
                    """
                    SELECT si.sku, p.name, p.unit, si.quantity
                    FROM sale_items si
                    JOIN products p ON p.sku = si.sku
                    WHERE si.sale_id = ?
                    ORDER BY si.id
                    """,
                    (sale["id"],),
                ).fetchall()
                row = dict(sale)
                row["items"] = [dict(item) for item in items]
                output.append(row)
        return output

    def get_report(self, from_date: str, to_date: str) -> list[dict[str, Any]]:
        try:
            start = datetime.strptime(from_date, "%Y-%m-%d")
            end = datetime.strptime(to_date, "%Y-%m-%d") + timedelta(days=1)
        except ValueError as exc:
            raise ValidationError("date must be YYYY-MM-DD") from exc
        if end <= start:
            raise ValidationError("to_date must be >= from_date")

        start_iso = start.replace(tzinfo=UTC).isoformat().replace("+00:00", "Z")
        end_iso = end.replace(tzinfo=UTC).isoformat().replace("+00:00", "Z")

        sql = """
        WITH sku_set AS (
            SELECT sku FROM products
            UNION
            SELECT sku FROM ledger
        )
        SELECT
            s.sku,
            COALESCE(p.name, '') AS name,
            COALESCE(p.unit, '') AS unit,
            COALESCE((
                SELECT SUM(l.change) FROM ledger l
                WHERE l.sku = s.sku AND l.created_at < ?
            ), 0) AS opening,
            COALESCE((
                SELECT SUM(l.change) FROM ledger l
                WHERE l.sku = s.sku AND l.change > 0
                AND l.created_at >= ? AND l.created_at < ?
            ), 0) AS imported,
            COALESCE((
                SELECT -SUM(l.change) FROM ledger l
                WHERE l.sku = s.sku AND l.change < 0
                AND l.created_at >= ? AND l.created_at < ?
            ), 0) AS exported
        FROM sku_set s
        LEFT JOIN products p ON p.sku = s.sku
        ORDER BY s.sku
        """
        with self.db.connection() as conn:
            rows = conn.execute(
                sql, (start_iso, start_iso, end_iso, start_iso, end_iso)
            ).fetchall()
        result: list[dict[str, Any]] = []
        for row in rows:
            opening = int(row["opening"])
            imported = int(row["imported"])
            exported = int(row["exported"])
            result.append(
                {
                    "sku": row["sku"],
                    "name": row["name"],
                    "unit": row["unit"],
                    "opening": opening,
                    "import": imported,
                    "export": exported,
                    "closing": opening + imported - exported,
                }
            )
        return result

    def validate_initial_import(self) -> None:
        with self.db.connection() as conn:
            row = conn.execute(
                "SELECT COUNT(*) AS c FROM products WHERE quantity > 0"
            ).fetchone()
        if int(row["c"]) > 0:
            raise ValidationError("Database must be empty or reset before import")

    def import_initial_stock_from_excel(
        self, file_path: str | Path, preview_only: bool = False
    ) -> list[dict[str, Any]]:
        rows = self._read_initial_stock_file(Path(file_path))
        if preview_only:
            return rows

        self.validate_initial_import()
        with self.db.transaction() as conn:
            for row in rows:
                sku = row["sku"]
                existing = conn.execute(
                    "SELECT sku FROM products WHERE sku = ?", (sku,)
                ).fetchone()
                if existing is None:
                    conn.execute(
                        "INSERT INTO products (sku, name, unit, quantity) VALUES (?, ?, ?, 0)",
                        (sku, row["name"], row["unit"]),
                    )
                else:
                    conn.execute(
                        "UPDATE products SET name = ?, unit = ? WHERE sku = ?",
                        (row["name"], row["unit"], sku),
                    )
                qty = int(row["quantity"])
                if qty > 0:
                    self._apply_stock_change(
                        conn=conn,
                        sku=sku,
                        change=qty,
                        movement_type="IMPORT",
                        ref_type="INIT",
                        ref_id=None,
                    )
        return rows

    def import_legacy_workbook(
        self, file_path: str | Path, preview_only: bool = False
    ) -> dict[str, Any]:
        stock_rows, txn_rows = self._read_legacy_workbook(Path(file_path))
        if preview_only:
            return {"initial_stock": stock_rows, "transactions": txn_rows}

        self.validate_initial_import()
        with self.db.transaction() as conn:
            for row in stock_rows:
                sku = row["sku"]
                existing = conn.execute(
                    "SELECT sku FROM products WHERE sku = ?", (sku,)
                ).fetchone()
                if existing is None:
                    conn.execute(
                        "INSERT INTO products (sku, name, unit, quantity) VALUES (?, ?, ?, 0)",
                        (sku, row["name"], row["unit"]),
                    )
                else:
                    conn.execute(
                        "UPDATE products SET name = ?, unit = ? WHERE sku = ?",
                        (row["name"], row["unit"], sku),
                    )
                qty = int(row["quantity"])
                if qty > 0:
                    self._apply_stock_change(
                        conn=conn,
                        sku=sku,
                        change=qty,
                        movement_type="IMPORT",
                        ref_type="INIT",
                        ref_id=None,
                    )

            for tx in txn_rows:
                change = int(tx["quantity"]) if tx["direction"] == "IMPORT" else -int(tx["quantity"])
                self._apply_stock_change(
                    conn=conn,
                    sku=tx["sku"],
                    change=change,
                    movement_type=tx["direction"],
                    ref_type="LEGACY_NHAP_LIEU",
                    ref_id=None,
                )
                conn.execute(
                    """
                    INSERT INTO stock_transactions (
                        transaction_code, employee_name, transaction_date, area, description,
                        sku, quantity, direction, note, bill_return, address, phone, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        tx["transaction_code"],
                        tx["employee_name"],
                        tx["transaction_date"],
                        tx["area"],
                        tx["description"],
                        tx["sku"],
                        tx["quantity"],
                        tx["direction"],
                        tx["note"],
                        tx["bill_return"],
                        tx["address"],
                        tx["phone"],
                        utc_now_iso(),
                    ),
                )
        return {"initial_stock": stock_rows, "transactions": txn_rows}

    def reset_database(self, mode: str = "full", confirm_token: str | None = None) -> None:
        if confirm_token != "RESET":
            raise ValidationError('Please type "RESET" to confirm')

        with self.db.transaction() as conn:
            conn.execute("DELETE FROM stock_transactions")
            conn.execute("DELETE FROM sale_items")
            conn.execute("DELETE FROM sales")
            conn.execute("DELETE FROM hold_items")
            conn.execute("DELETE FROM hold_sessions")
            conn.execute("DELETE FROM ledger")
            if mode == "full":
                conn.execute("DELETE FROM products")
            elif mode == "safe":
                conn.execute("UPDATE products SET quantity = 0")
            else:
                raise ValidationError("mode must be full or safe")

    def get_stock_transactions(self, limit: int = 500) -> list[dict[str, Any]]:
        with self.db.connection() as conn:
            rows = conn.execute(
                """
                SELECT transaction_code, employee_name, transaction_date, area, description,
                       sku, quantity, direction, note, bill_return, address, phone
                FROM stock_transactions
                ORDER BY id DESC
                LIMIT ?
                """,
                (limit,),
            ).fetchall()
        return [dict(row) for row in rows]

    def backup_db(self, dest_dir: str | Path) -> str:
        dest = Path(dest_dir)
        dest.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = dest / f"warehouse_backup_{stamp}.db"
        shutil.copy2(self.db.db_path, backup_path)
        return str(backup_path)

    def restore_db(self, backup_file: str | Path) -> None:
        source = Path(backup_file)
        if not source.exists() or source.suffix != ".db":
            raise ValidationError("Invalid backup file")
        shutil.copy2(source, self.db.db_path)

    def check_stock_consistency(self) -> list[dict[str, Any]]:
        sql = """
        WITH ledger_sum AS (
            SELECT sku, COALESCE(SUM(change), 0) AS ledger_qty
            FROM ledger
            GROUP BY sku
        )
        SELECT p.sku, p.quantity AS product_qty, COALESCE(l.ledger_qty, 0) AS ledger_qty
        FROM products p
        LEFT JOIN ledger_sum l ON l.sku = p.sku
        WHERE p.quantity <> COALESCE(l.ledger_qty, 0)
        ORDER BY p.sku
        """
        with self.db.connection() as conn:
            rows = conn.execute(sql).fetchall()
        return [dict(row) for row in rows]

    def _read_initial_stock_file(self, file_path: Path) -> list[dict[str, Any]]:
        if not file_path.exists():
            raise ValidationError("File not found")
        if file_path.suffix.lower() in {".csv", ".txt"}:
            return self._parse_csv(file_path)

        if file_path.suffix.lower() in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise ValidationError(
                    "openpyxl is required for .xlsx files. Install optional dependency: pip install .[excel]"
                ) from exc

            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
            return self._parse_tabular_rows(rows)

        raise ValidationError("Only .xlsx or .csv files are supported")

    def _read_legacy_workbook(
        self, file_path: Path
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        if not file_path.exists():
            raise ValidationError("File not found")
        if file_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValidationError("Legacy import supports .xlsx/.xlsm only")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValidationError(
                "openpyxl is required for legacy workbook import. Install optional dependency: pip install .[excel]"
            ) from exc

        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            initial_ws = wb["SỐ DƯ ĐẦU KỲ"]
            tx_ws = wb["NHAP LIEU"]
        except KeyError as exc:
            wb.close()
            raise ValidationError("Workbook must contain sheets: SỐ DƯ ĐẦU KỲ and NHAP LIEU") from exc

        initial_rows = [list(r) for r in initial_ws.iter_rows(values_only=True)]
        tx_rows = [list(r) for r in tx_ws.iter_rows(values_only=True)]
        wb.close()
        return self._parse_legacy_initial_rows(initial_rows), self._parse_legacy_transaction_rows(tx_rows)

    def _parse_csv(self, file_path: Path) -> list[dict[str, Any]]:
        with file_path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            rows = list(reader)
        return self._parse_tabular_rows(rows)

    def _parse_tabular_rows(self, rows: list[Any]) -> list[dict[str, Any]]:
        if not rows:
            raise ValidationError("File is empty")

        header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        idx = self._resolve_initial_import_header(header)

        parsed: list[dict[str, Any]] = []
        seen_sku: set[str] = set()
        for row in rows[1:]:
            if row is None:
                continue
            values = list(row)
            sku = self._safe_cell(values, idx["sku"]).strip()
            if not sku:
                continue
            name = self._safe_cell(values, idx["name"]).strip()
            unit = self._safe_cell(values, idx["unit"]).strip()
            qty_str = self._safe_cell(values, idx["quantity"]).strip()
            if not name or not unit:
                raise ValidationError(f"Missing name/unit for SKU: {sku}")
            try:
                quantity = int(float(qty_str))
            except ValueError as exc:
                raise ValidationError(f"Invalid quantity for SKU: {sku}") from exc
            if quantity < 0:
                raise ValidationError(f"Quantity must be >= 0 for SKU: {sku}")
            if sku in seen_sku:
                raise ValidationError(f"Duplicate SKU in file: {sku}")
            seen_sku.add(sku)
            parsed.append(
                {"sku": sku, "name": name, "unit": unit, "quantity": quantity}
            )

        if not parsed:
            raise ValidationError("No valid rows found")
        return parsed

    def _parse_legacy_initial_rows(self, rows: list[list[Any]]) -> list[dict[str, Any]]:
        if not rows:
            raise ValidationError("Sheet SỐ DƯ ĐẦU KỲ is empty")

        header_idx, header_map = self._find_legacy_header_map(
            rows=rows,
            required={
                "sku": ("mã hàng", "ma hang", "sku"),
                "name": ("tên hàng hóa", "ten hang hoa", "name"),
                "unit": ("đvt", "dvt", "unit"),
                "quantity": ("tồn", "ton", "quantity", "qty"),
            },
            sheet_name="SỐ DƯ ĐẦU KỲ",
        )
        parsed: list[dict[str, Any]] = []
        seen_sku: set[str] = set()
        for row in rows[header_idx + 1 :]:
            sku = self._safe_cell(row, header_map["sku"]).strip()
            name = self._safe_cell(row, header_map["name"]).strip()
            unit = self._safe_cell(row, header_map["unit"]).strip()
            qty_str = self._safe_cell(row, header_map["quantity"]).strip()
            if not sku:
                continue
            if not name or not unit:
                raise ValidationError(f"Missing name/unit for SKU: {sku}")
            try:
                quantity = int(float(qty_str or "0"))
            except ValueError as exc:
                raise ValidationError(f"Invalid quantity for SKU: {sku}") from exc
            if quantity < 0:
                raise ValidationError(f"Quantity must be >= 0 for SKU: {sku}")
            if sku in seen_sku:
                raise ValidationError(f"Duplicate SKU in SỐ DƯ ĐẦU KỲ: {sku}")
            seen_sku.add(sku)
            parsed.append(
                {"sku": sku, "name": name, "unit": unit, "quantity": quantity}
            )
        if not parsed:
            raise ValidationError("No valid inventory rows in SỐ DƯ ĐẦU KỲ")
        return parsed

    def _parse_legacy_transaction_rows(self, rows: list[list[Any]]) -> list[dict[str, Any]]:
        if len(rows) < 2:
            return []
        header_idx, header_map = self._find_legacy_header_map(
            rows=rows,
            required={
                "employee_name": ("nv đề xuất", "nhân viên đề xuất", "nhan vien de xuat"),
                "transaction_date": ("ngày", "ngay", "date"),
                "area": ("khu vực", "khu vuc", "area"),
                "description": ("diễn giải", "dien giai", "description"),
                "import_sku": ("nhập", "nhap", "import"),
                "export_sku": ("xuất", "xuat", "export"),
                "quantity": ("soluong", "số lượng", "so luong", "quantity", "qty"),
            },
            sheet_name="NHAP LIEU",
        )

        note_idx = self._optional_header_index(
            rows[header_idx], ("ghi chú", "ghi chu", "note")
        )
        bill_idx = self._optional_header_index(
            rows[header_idx], ("trả bill", "tra bill", "bill")
        )
        address_idx = self._optional_header_index(
            rows[header_idx], ("địa chỉ", "dia chi", "address")
        )
        phone_idx = self._optional_header_index(
            rows[header_idx], ("điện thoại", "dien thoai", "phone", "tel")
        )

        parsed: list[dict[str, Any]] = []
        for idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            sku_import = self._safe_cell(row, header_map["import_sku"]).strip()
            sku_export = self._safe_cell(row, header_map["export_sku"]).strip()
            qty_raw = self._safe_cell(row, header_map["quantity"]).strip()
            if not sku_import and not sku_export:
                continue
            try:
                quantity = int(float(qty_raw))
            except ValueError as exc:
                raise ValidationError(f"Invalid quantity at NHAP LIEU row {idx}") from exc
            if quantity <= 0:
                continue
            direction = "IMPORT" if sku_import else "EXPORT"
            sku = sku_import if sku_import else sku_export
            parsed.append(
                {
                    "transaction_code": f"LEGACY-{idx:05d}",
                    "employee_name": self._safe_cell(row, header_map["employee_name"]).strip(),
                    "transaction_date": self._normalize_legacy_date(
                        self._safe_cell(row, header_map["transaction_date"]).strip()
                    ),
                    "area": self._safe_cell(row, header_map["area"]).strip(),
                    "description": self._safe_cell(row, header_map["description"]).strip(),
                    "sku": sku,
                    "quantity": quantity,
                    "direction": direction,
                    "note": self._safe_cell(row, note_idx).strip() if note_idx is not None else "",
                    "bill_return": self._safe_cell(row, bill_idx).strip() if bill_idx is not None else "",
                    "address": self._safe_cell(row, address_idx).strip() if address_idx is not None else "",
                    "phone": self._safe_cell(row, phone_idx).strip() if phone_idx is not None else "",
                }
            )
        return parsed

    def _resolve_initial_import_header(self, header: list[str]) -> dict[str, int]:
        normalized = [h.casefold() for h in header]
        mapping: dict[str, tuple[str, ...]] = {
            "sku": ("mã hàng", "ma hang", "sku", "mã sản phẩm", "ma san pham"),
            "name": ("tên hàng hóa", "ten hang hoa", "name", "tên", "ten"),
            "unit": ("đvt", "dvt", "unit"),
            "quantity": ("tồn", "ton", "quantity", "qty"),
        }
        resolved: dict[str, int] = {}
        for key, aliases in mapping.items():
            found = None
            for alias in aliases:
                if alias in normalized:
                    found = normalized.index(alias)
                    break
            if found is None:
                raise ValidationError(f"Missing required column: {key}")
            resolved[key] = found
        return resolved

    @staticmethod
    def _safe_cell(row: list[Any], index: int) -> str:
        if index is None:
            return ""
        if index >= len(row):
            return ""
        value = row[index]
        return "" if value is None else str(value)

    def _find_legacy_header_map(
        self,
        rows: list[list[Any]],
        required: dict[str, tuple[str, ...]],
        sheet_name: str,
    ) -> tuple[int, dict[str, int]]:
        for row_idx, row in enumerate(rows):
            header_map: dict[str, int] = {}
            for key, aliases in required.items():
                idx = self._optional_header_index(row, aliases)
                if idx is None:
                    break
                header_map[key] = idx
            if len(header_map) == len(required):
                return row_idx, header_map
        missing = ", ".join(required.keys())
        raise ValidationError(f"Missing required headers in {sheet_name}: {missing}")

    @staticmethod
    def _optional_header_index(row: list[Any], aliases: tuple[str, ...]) -> int | None:
        normalized = [str(c).strip().casefold() if c is not None else "" for c in row]
        for alias in aliases:
            alias_norm = alias.casefold()
            if alias_norm in normalized:
                return normalized.index(alias_norm)
        return None

    @staticmethod
    def _normalize_legacy_date(raw: str) -> str:
        if not raw:
            return ""
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
        # Legacy sheet commonly stores day/month without year. Assume current year.
        sep = "/" if "/" in raw else "-" if "-" in raw else None
        if sep:
            parts = [p.strip() for p in raw.split(sep)]
            if len(parts) == 2 and all(p.isdigit() for p in parts):
                day = int(parts[0])
                month = int(parts[1])
                year = datetime.now().year
                try:
                    return datetime(year=year, month=month, day=day).strftime("%Y-%m-%d")
                except ValueError:
                    return raw
        return raw

    @staticmethod
    def _validate_positive_qty(quantity: int) -> None:
        if int(quantity) <= 0:
            raise ValidationError("Quantity must be > 0")

    def _normalize_hold_items(self, items: list[dict[str, Any]]) -> list[HoldItemInput]:
        if not items:
            raise ValidationError("Hold items are required")
        merged: dict[str, int] = {}
        for item in items:
            sku = str(item.get("sku", "")).strip()
            quantity = int(item.get("quantity", 0))
            if not sku:
                raise ValidationError("SKU is required for hold")
            self._validate_positive_qty(quantity)
            merged[sku] = merged.get(sku, 0) + quantity
        return [HoldItemInput(sku=sku, quantity=qty) for sku, qty in merged.items()]

    def _apply_stock_change(
        self,
        conn: sqlite3.Connection,
        sku: str,
        change: int,
        movement_type: str,
        ref_type: str,
        ref_id: int | None,
    ) -> None:
        sku = sku.strip()
        product = conn.execute(
            "SELECT quantity FROM products WHERE sku = ?", (sku,)
        ).fetchone()
        if product is None:
            raise NotFoundError(f"Product not found: {sku}")

        current_qty = int(product["quantity"])
        new_qty = current_qty + change
        if new_qty < 0:
            raise InsufficientStockError(
                f"Insufficient stock for {sku}: have {current_qty}, need {abs(change)}"
            )

        conn.execute("UPDATE products SET quantity = ? WHERE sku = ?", (new_qty, sku))
        conn.execute(
            """
            INSERT INTO ledger (sku, change, type, ref_type, ref_id, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (sku, change, movement_type, ref_type, ref_id, utc_now_iso()),
        )

    @staticmethod
    def _new_code(prefix: str) -> str:
        return f"{prefix}-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid4().hex[:6].upper()}"
